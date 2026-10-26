"""
End-to-end API tests (the environment and a fresh test database come from conftest.py).
Set FOX_FILE=/path/to/case.fox to also run the round-trip on a real case.
"""
import json
import os
import threading
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from app.main import app
from app import cases as C

ADMIN_PW = os.environ["ADMIN_PASSWORD"]

FIXTURE = {
    "storage_format_version": 7, "locked": True, "case_id": "T-1", "client": "ACME", "start_date": "2026-01-01",
    "summary": "test",
    "timeline": [{"recid": 1, "w2ui": {}, "date_time": "2026-01-01 10:00:00", "event_type": "EventLog",
                  "event_host": "host-a", "event_source_host": "evil.example", "direction": "->",
                  "killchain": "C2", "event_data": "beacon", "notes": "", "visual": True, "followup": False,
                  "owner": "jd"}],
    "investigated_systems": [{"recid": 1, "date_added": 1791220203619, "date_added_": "2026-10-05T17:10:03.619Z",
                              "date_updated": 1791220203619, "w2ui": {}, "hostname": "host-a", "summary": "x",
                              "analysis_required": True, "analyst": "jd", "analysis_status": "Assigned"}],
    "malware": [{"recid": 1, "date_added": 1, "w2ui": {}, "text": "bad.exe", "hostname": "host-a",
                 "md5": "d41d8cd98f00b204e9800998ecf8427e"}],
    "compromised_accounts": [], "network_indicators": [
        {"recid": 1, "date_added": 1, "w2ui": {}, "ip": "10.0.0.1", "port": "443", "domainname": "evil.example",
         "context": "", "malware": "bad.exe"}],
    "exfiltration": [], "hosts": [],
    # Records are listed in each section's default order (systems by hostname), so this file round-trips exactly.
    "systems": [{"recid": 2, "date_added": 1, "text": "evil.example", "system_type": "Attacker Infra", "pin": True},
                {"recid": 1, "date_added": 1, "text": "host-a", "ip": "10.0.0.5", "system_type": "Desktop",
                 "os": "", "owner": "", "description": "", "pin": "", "w2ui": {}}],
    "osint": [], "investigators": [{"recid": 1, "w2ui": {}, "text": "jd", "full_name": "Jane Doe",
                                    "affiliation": "", "role": "IR"}],
    "evidence": [], "actions": [], "casenotes": [],
    **{k: v for k, v in C.S.default_lookups().items()},
}


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as c:
        yield c


def login(c, user="admin", pw=ADMIN_PW):
    r = c.post("/api/login", json={"username": user, "password": pw})
    assert r.status_code == 200, r.text
    c.headers["X-CSRF-Token"] = r.json()["csrf"]
    return c


@pytest.fixture(scope="module")
def admin(client):
    return login(client)


@pytest.fixture(scope="module")
def case(admin):
    r = admin.post("/api/cases/import", content=json.dumps(FIXTURE))
    assert r.status_code == 200, r.text
    return r.json()["id"]


def sec(c, case, s):
    return c.get(f"/api/cases/{case}/sections/{s}").json()["records"]


def test_security_headers_and_auth(client):
    r = client.get("/")
    assert "default-src 'self'" in r.headers["content-security-policy"]
    assert r.headers["x-frame-options"] == "DENY"
    fresh = TestClient(app)
    assert fresh.get("/api/cases").status_code == 401
    assert fresh.post("/api/login", json={"username": "admin", "password": "wrong-password!"}).status_code == 401
    assert fresh.get("/docs").status_code in (404, 405) or "swagger" not in fresh.get("/docs").text.lower()


def test_csrf_required(admin, case):
    tok = admin.headers.pop("X-CSRF-Token")
    try:
        r = admin.post(f"/api/cases/{case}/sections/osint", json={"data": {"url": "https://x"}})
        assert r.status_code == 403
    finally:
        admin.headers["X-CSRF-Token"] = tok


def test_roundtrip_identical_fixture(admin, case):
    r = admin.get(f"/api/cases/{case}/export?locked=true")
    assert r.status_code == 200
    assert r.content == json.dumps(FIXTURE, ensure_ascii=False, indent="\t").encode()


def _canonical(doc):
    """A .fox document with every section in recid order, text joined into single lines (as imported)."""
    def fix(v):
        return C.S.join_lines(v) if isinstance(v, str) else v
    out = {}
    for k, v in doc.items():
        if k in C.S.SECTIONS and isinstance(v, list):
            out[k] = sorted(({kk: fix(vv) for kk, vv in r.items()} for r in v), key=lambda r: r["recid"])
        else:
            out[k] = fix(v) if k in C.S.CASE_FIELDS else v
    return out


@pytest.mark.skipif(not os.environ.get("FOX_FILE"), reason="FOX_FILE not set")
def test_roundtrip_real_file(admin):
    """Exports re-order records into the default order; the content is unchanged and a second round trip
    is byte-identical."""
    raw = Path(os.environ["FOX_FILE"]).read_bytes()
    pk = admin.post("/api/cases/import", content=raw).json()["id"]
    out = admin.get(f"/api/cases/{pk}/export?locked=true").content
    assert _canonical(json.loads(out)) == _canonical(json.loads(raw))
    pk2 = admin.post("/api/cases/import", content=out).json()["id"]
    assert admin.get(f"/api/cases/{pk2}/export?locked=true").content == out


def test_export_unlocked_by_default(admin, case):
    doc = json.loads(admin.get(f"/api/cases/{case}/export").content)
    assert doc["locked"] is False


def test_reference_must_be_declared(admin, case):
    r = admin.post(f"/api/cases/{case}/sections/timeline",
                   json={"data": {"event_host": "unknown-host", "event_data": "x", "date_time": "2026-01-02 00:00"}})
    assert r.status_code == 422
    assert r.json()["missing_refs"][0] == {"field": "event_host", "section": "systems", "value": "unknown-host"}
    r = admin.post(f"/api/cases/{case}/sections/timeline", json={"data": {"owner": "nobody"}})
    assert r.status_code == 422
    r = admin.post(f"/api/cases/{case}/sections/timeline", json={"data": {"event_type": "NotAType"}})
    assert r.status_code == 422


def test_create_and_formats(admin, case):
    r = admin.post(f"/api/cases/{case}/sections/timeline",
                   json={"data": {"event_host": "host-a", "date_time": "2026-01-02T03:04", "event_data": "<script>"}})
    assert r.status_code == 200, r.text
    d = r.json()["data"]
    assert d["date_time"] == "2026-01-02 03:04:00"
    assert d["event_data"] == "<script>"            # stored verbatim, rendered as text by the UI
    assert r.json()["recid"] == 2
    r = admin.post(f"/api/cases/{case}/sections/casenotes", json={"data": {"note": "hello", "owner": "jd"}})
    d = r.json()["data"]
    assert isinstance(d["date_added"], int) and d["date_added_"].endswith("Z")


@pytest.mark.parametrize("section,data", [
    ("systems", {"text": "x1", "ip": "999.1.1.1"}),
    ("network_indicators", {"port": "70000"}),
    ("malware", {"text": "a", "md5": "nothex"}),
    ("timeline", {"date_time": "2026-13-01 00:00:00"}),
    ("timeline", {"notes": "bad\x00null"}),
    ("systems", {"text": "line\nbreak"}),
    ("systems", {"text": "y", "unknown_field": 1}),
    ("timeline", {"visual": "maybe"}),
    ("investigators", {"text": "x" * 21}),
])
def test_validation_rejects(admin, case, section, data):
    r = admin.post(f"/api/cases/{case}/sections/{section}", json={"data": data})
    assert r.status_code == 422, r.text


def test_unique_system_names(admin, case):
    r = admin.post(f"/api/cases/{case}/sections/systems", json={"data": {"text": "HOST-A"}})
    assert r.status_code == 422


def test_field_level_merge_and_conflict(admin, case):
    rec = next(r for r in sec(admin, case, "systems") if r["data"]["text"] == "host-a")
    url = f"/api/cases/{case}/sections/systems/{rec['recid']}"
    # user 1 edits os, user 2 edits owner, both from the same base -> both succeed
    r1 = admin.patch(url, json={"changes": {"os": "Win10"}, "expected": {"os": rec["data"]["os"]}})
    r2 = admin.patch(url, json={"changes": {"owner": "Bob"}, "expected": {"owner": rec["data"]["owner"]}})
    assert r1.status_code == 200 and r2.status_code == 200
    d = r2.json()["data"]
    assert d["os"] == "Win10" and d["owner"] == "Bob"
    # user 3 edits os from the stale base -> conflict
    r3 = admin.patch(url, json={"changes": {"os": "Linux"}, "expected": {"os": ""}})
    assert r3.status_code == 409
    assert r3.json()["conflicts"][0]["current"] == "Win10"


def test_rename_cascades(admin, case):
    systems = sec(admin, case, "systems")
    host = next(s for s in systems if s["data"]["text"] == "host-a")
    r = admin.patch(f"/api/cases/{case}/sections/systems/{host['recid']}",
                    json={"changes": {"text": "host-a.corp"}, "expected": {"text": "host-a"}})
    assert r.status_code == 200
    assert all(t["data"].get("event_host") != "host-a" for t in sec(admin, case, "timeline"))
    assert sec(admin, case, "investigated_systems")[0]["data"]["hostname"] == "host-a.corp"
    assert sec(admin, case, "malware")[0]["data"]["hostname"] == "host-a.corp"
    inv = sec(admin, case, "investigators")[0]
    admin.patch(f"/api/cases/{case}/sections/investigators/{inv['recid']}",
                json={"changes": {"text": "jdoe"}, "expected": {"text": "jd"}})
    assert next(t for t in sec(admin, case, "timeline") if t["recid"] == 1)["data"]["owner"] == "jdoe"


def test_delete_referenced_blocked(admin, case):
    host = next(s for s in sec(admin, case, "systems") if s["data"]["text"] == "host-a.corp")
    r = admin.post(f"/api/cases/{case}/sections/systems/{host['recid']}/delete", json={"version": host["version"]})
    assert r.status_code == 409 and r.json()["references"]
    stale = admin.post(f"/api/cases/{case}/sections/timeline/1/delete", json={"version": 0})
    assert stale.status_code == 409


def test_viewer_is_read_only(admin, case):
    r = admin.post("/api/users", json={"username": "viewer1", "password": "viewer-password-1", "role": "user",
                                            "must_change_password": False})
    assert r.status_code == 200, r.text
    uid = r.json()["id"]
    assert admin.post(f"/api/cases/{case}/members", json={"user_id": uid, "permission": "viewer"}).status_code == 200
    v = login(TestClient(app), "viewer1", "viewer-password-1")
    assert v.get(f"/api/cases/{case}/sections/timeline").status_code == 200
    assert v.post(f"/api/cases/{case}/sections/osint", json={"data": {}}).status_code == 403
    assert v.get(f"/api/cases/{case}/export").status_code == 200
    assert v.get("/api/users").status_code == 403


def test_concurrent_writers(admin, case):
    errors = []

    def worker(n):
        c = login(TestClient(app))
        for i in range(10):
            r = c.post(f"/api/cases/{case}/sections/casenotes", json={"data": {"note": f"w{n}-{i}"}})
            if r.status_code != 200:
                errors.append(r.text)

    threads = [threading.Thread(target=worker, args=(n,)) for n in range(6)]
    [t.start() for t in threads]
    [t.join() for t in threads]
    assert not errors
    notes = sec(admin, case, "casenotes")
    assert len({n["recid"] for n in notes}) == len(notes)
    doc = json.loads(admin.get(f"/api/cases/{case}/export").content)
    assert not C.integrity_errors(doc)


def test_import_rejects_bad_refs_unless_declared(admin):
    bad = json.loads(json.dumps(FIXTURE))
    bad["timeline"][0]["event_host"] = "ghost"
    r = admin.post("/api/cases/import", content=json.dumps(bad))
    assert r.status_code == 422
    r = admin.post("/api/cases/import?declare_missing=true", content=json.dumps(bad))
    assert r.status_code == 200 and any("ghost" in w for w in r.json()["warnings"])
    assert admin.post("/api/cases/import", content=b"{not json").status_code == 422
    assert admin.post("/api/cases/import", content=json.dumps({"storage_format_version": 99})).status_code == 422


def test_import_rejects_deeply_nested_json(admin):
    """A crafted .fox with a deeply nested unknown key must be refused, not stored (it would otherwise break
    .fox export: json.dumps recurses over the passthrough and 500s). A normal .fox still imports."""
    bomb = '{"storage_format_version": 7, "case_id": "BOMB", "x": ' + "[" * 500 + "]" * 500 + "}"
    r = admin.post("/api/cases/import", content=bomb)
    assert r.status_code == 422 and "deep" in r.text.lower(), r.text
    # a legitimate unknown top-level key (shallow forward-compat) is still carried through
    doc = json.loads(json.dumps(FIXTURE)); doc["case_id"] = "PASS1"; doc["future_field"] = {"a": [1, 2, 3]}
    r = admin.post("/api/cases/import", content=json.dumps(doc))
    assert r.status_code == 200, r.text
    cid = r.json()["id"]
    assert admin.get(f"/api/cases/{cid}/export").status_code == 200       # .fox still serialises


def test_body_limit(admin, case):
    r = admin.post(f"/api/cases/{case}/sections/casenotes", content=b"x" * (2 * 1024 * 1024),
                   headers={"content-type": "application/json"})
    assert r.status_code == 413


def test_case_meta_update(admin, case):
    meta = admin.get(f"/api/cases/{case}").json()["meta"]
    r = admin.patch(f"/api/cases/{case}", json={"changes": {"client": "New"}, "expected": {"client": meta["client"]}})
    assert r.status_code == 200
    r = admin.patch(f"/api/cases/{case}", json={"changes": {"client": "Other"}, "expected": {"client": "stale"}})
    assert r.status_code == 409


def test_report_data_is_read_only(admin, case):
    before_rev = admin.get(f"/api/cases/{case}/revision").json()["revision"]
    before_audit = admin.get(f"/api/cases/{case}/audit?limit=1").json()[0]["id"]
    r = admin.get(f"/api/cases/{case}/report-data")
    assert r.status_code == 200
    body = r.json()
    assert set(body["sections"]) >= {"timeline", "systems", "investigated_systems"}
    assert body["meta"].get("vtapikey", "") in ("", "********")     # secrets never in reports
    assert admin.get(f"/api/cases/{case}/revision").json()["revision"] == before_rev
    assert admin.get(f"/api/cases/{case}/audit?limit=1").json()[0]["id"] == before_audit
    assert admin.post(f"/api/cases/{case}/report-data", json={}).status_code in (404, 405)


@pytest.mark.parametrize("given,stored", [
    ("2020-02-21 14:20:37", "2020-02-21 14:20:37"),          # no zone -> taken as UTC
    ("2020-02-21T14:20:37Z", "2020-02-21 14:20:37"),
    ("2020-02-21 14:20:37 UTC", "2020-02-21 14:20:37"),
    ("2020-02-21T15:20:37+01:00", "2020-02-21 14:20:37"),    # offsets converted to UTC
    ("2020-12-31 23:30:00-01:00", "2021-01-01 00:30:00"),
])
def test_datetimes_stored_in_utc(admin, case, given, stored):
    r = admin.post(f"/api/cases/{case}/sections/casenotes", json={"data": {"note": "x"}})
    assert r.json()["data"]["date_added_"].endswith("Z")
    r = admin.post(f"/api/cases/{case}/sections/exfiltration", json={"data": {"exfiltime": given}})
    assert r.status_code == 200, r.text
    assert r.json()["data"]["exfiltime"] == stored


def test_case_info_includes_section_counts(admin, case):
    counts = admin.get(f"/api/cases/{case}").json()["counts"]
    for section in ("timeline", "systems", "investigators", "evidence", "actions", "casenotes", "osint"):
        assert counts[section] == len(sec(admin, case, section))


def test_excel_export(admin, case):
    import io, zipfile
    from openpyxl import load_workbook
    evil = "=cmd|' /C calc'!A0"
    admin.post(f"/api/cases/{case}/sections/casenotes", json={"data": {"note": evil}})
    admin.post(f"/api/cases/{case}/sections/osint", json={"data": {"url": "=HYPERLINK(\"http://x\")", "desc": "+1-1"}})
    r = admin.get(f"/api/cases/{case}/export.xlsx")
    assert r.status_code == 200 and r.headers["content-disposition"].endswith('.xlsx"')
    raw = r.content
    with zipfile.ZipFile(io.BytesIO(raw)) as z:          # no formulas anywhere in the workbook
        assert not any(b"<f>" in z.read(n) or b"<f " in z.read(n) for n in z.namelist() if n.startswith("xl/worksheets/"))
    wb = load_workbook(io.BytesIO(raw))
    assert wb.sheetnames == ["Cover", "Report", "Timeline", "Investigated Systems", "Malware-Tools", "Compromised Accounts",
                             "Network Indicators", "Exfiltration", "OSInt", "Systems", "Investigators", "Evidence",
                             "Action Items", "Case Notes"]
    fox = json.loads(admin.get(f"/api/cases/{case}/export").content)
    for sheet, key in [("Timeline", "timeline"), ("Systems", "systems"), ("Case Notes", "casenotes"), ("OSInt", "osint")]:
        assert wb[sheet].max_row - 1 == len(fox[key])
    tl = wb["Timeline"]
    assert tl["B1"].value == "Date/Time (UTC)"
    first = tl["B2"].value
    assert first.strftime("%Y-%m-%d %H:%M:%S") == fox["timeline"][0]["date_time"]     # real datetime cell
    notes = [c.value for c in wb["Case Notes"]["C"]]
    assert evil in notes                                                              # kept verbatim as text
    cell = next(c for c in wb["Case Notes"]["C"] if c.value == evil)
    assert cell.data_type == "s"
    assert "vtapikey" not in str([c.value for row in wb["Cover"].iter_rows() for c in row])
    assert "Source" not in [c.value for c in wb["Cover"]["B"]]
    assert "export_xlsx" in [a["action"] for a in admin.get(f"/api/cases/{case}/audit?limit=3").json()]


def test_ui_files_revalidated_and_no_password_inputs_for_secrets(client):
    r = client.get("/app.js")
    assert r.headers["cache-control"] == "no-cache" and r.headers.get("etag")
    assert client.get("/app.js", headers={"If-None-Match": r.headers["etag"]}).status_code == 304
    js = r.text
    assert 'input.type = "password"' not in js            # secrets are masked text, not password fields
    assert 'type: "password"' not in client.get("/apiexplorer.js").text


def test_max_records_per_case(admin, monkeypatch):
    """MAX_RECORDS_PER_CASE caps total records across all sections (create, bulk, import); 0 = unlimited."""
    from app.config import settings
    r = admin.post("/api/cases", json={"case_id": "CAP-1", "client": "x"})
    cid = r.json()["id"]
    admin.post(f"/api/cases/{cid}/sections/systems", json={"data": {"text": "h"}})   # 1 record
    monkeypatch.setattr(settings, "max_records_per_case", 3)
    # two more -> ok (total 3); the fourth -> refused
    assert admin.post(f"/api/cases/{cid}/sections/casenotes", json={"data": {"note": "a"}}).status_code == 200
    assert admin.post(f"/api/cases/{cid}/sections/casenotes", json={"data": {"note": "b"}}).status_code == 200
    r = admin.post(f"/api/cases/{cid}/sections/casenotes", json={"data": {"note": "c"}})
    assert r.status_code == 422 and "MAX_RECORDS_PER_CASE" in r.text
    # bulk respects the cap and stays all-or-nothing (count unchanged)
    before = sum(admin.get(f"/api/cases/{cid}").json()["counts"].values())
    r = admin.post(f"/api/cases/{cid}/sections/casenotes/bulk", json={"records": [{"note": "x"}, {"note": "y"}]})
    assert r.status_code == 422
    assert sum(admin.get(f"/api/cases/{cid}").json()["counts"].values()) == before
    # import of a case exceeding the cap is refused
    big = json.loads(json.dumps(FIXTURE)); big["case_id"] = "CAP-IMP"
    big["casenotes"] = [{"recid": i + 1, "note": str(i), "w2ui": {}} for i in range(5)]
    assert admin.post("/api/cases/import", content=json.dumps(big)).status_code == 422
    # 0 = unlimited
    monkeypatch.setattr(settings, "max_records_per_case", 0)
    assert admin.post(f"/api/cases/{cid}/sections/casenotes", json={"data": {"note": "d"}}).status_code == 200
    admin.post(f"/api/cases/{cid}/delete", json={"confirm": "CAP-1"})


def test_export_rate_limit(admin, case, monkeypatch):
    """Per-principal rate limit spans all export formats; 0 = unlimited."""
    from app.config import settings
    from app.db import db
    with db() as c:
        c.execute("DELETE FROM login_failures WHERE k LIKE 'rl:export:%'")
    monkeypatch.setattr(settings, "export_rate_per_min", 2)
    assert admin.get(f"/api/cases/{case}/export").status_code == 200
    assert admin.get(f"/api/cases/{case}/export.xlsx").status_code == 200
    r = admin.get(f"/api/cases/{case}/export.pdf")          # 3rd export in the window -> 429
    assert r.status_code == 429 and "slow down" in r.text
    monkeypatch.setattr(settings, "export_rate_per_min", 0)
    assert admin.get(f"/api/cases/{case}/export").status_code == 200


def test_rate_limit_helper():
    """The sliding-window limiter (also used for on-demand backups): limit enforced, 0 = unlimited."""
    import pytest
    from fastapi import HTTPException
    from app import auth
    from app.db import db
    with db() as c:
        c.execute("DELETE FROM login_failures WHERE k LIKE 'rl:unit:%'")
    auth.rate_limit("unit", "x", 2, 60)
    auth.rate_limit("unit", "x", 2, 60)
    with pytest.raises(HTTPException) as e:
        auth.rate_limit("unit", "x", 2, 60)
    assert e.value.status_code == 429
    auth.rate_limit("unit", "x", 0, 60)          # 0 = unlimited: no raise
    auth.rate_limit("unit", "other", 2, 60)      # a different principal has its own window
