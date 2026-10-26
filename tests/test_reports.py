"""Report page (API, activity log, never in the .fox) and the Excel / PDF / LaTeX exports."""
import io
import json
import re

import pytest
from fastapi.testclient import TestClient

from tests.test_app import FIXTURE, login
from app.main import app

EVIL = r"\input{/etc/passwd} \write18{id} $x^2$ 50% & #1 _u_ ~ <img src=/etc/passwd> <a href='x'>y</a> =cmd|' /C calc'!A0"


@pytest.fixture(scope="module")
def admin():
    with TestClient(app) as c:
        yield login(c)


@pytest.fixture(scope="module")
def case(admin):
    fx = json.loads(json.dumps(FIXTURE))
    fx["case_id"] = "REP-1"
    fx["timeline"][0]["visual"] = True
    fx["timeline"][0]["event_data"] = EVIL
    r = admin.post("/api/cases/import", content=json.dumps(fx))
    assert r.status_code == 200, r.text
    return r.json()["id"]


def test_report_get_set_and_audit(admin, case):
    r = admin.get(f"/api/cases/{case}/report").json()
    assert r["executive_summary"] == "" and r["context"] == "test" and r["updated_by"] == ""
    body = {"executive_summary": "Ransomware intrusion.\n\nContained on day 3.",
            "context": "Client called on 2026-01-01.\nScope: 3 sites.", "conclusions": "Initial access via VPN.",
            "lessons_learned": "- MFA on VPN\n- EDR everywhere"}
    r = admin.patch(f"/api/cases/{case}/report", json={"changes": body})
    assert r.status_code == 200, r.text
    got = admin.get(f"/api/cases/{case}/report").json()
    assert {k: got[k] for k in body} == body and got["updated_by"] == "admin"
    # context is the case Summary (multi-line in the app, single line in the .fox)
    assert admin.get(f"/api/cases/{case}").json()["meta"]["summary"] == body["context"]
    log = admin.get(f"/api/cases/{case}/audit?limit=1").json()[0]
    assert log["action"] == "update_report" and log["after"]["conclusions"] == "Initial access via VPN."


def test_report_conflict_and_validation(admin, case):
    cur = admin.get(f"/api/cases/{case}/report").json()
    r = admin.patch(f"/api/cases/{case}/report", json={"changes": {"conclusions": "new"},
                                                       "expected": {"conclusions": "stale"}})
    assert r.status_code == 409 and r.json()["conflicts"][0]["current"] == cur["conclusions"]
    assert admin.patch(f"/api/cases/{case}/report", json={"changes": {"nope": "x"}}).status_code == 422
    assert admin.patch(f"/api/cases/{case}/report", json={"changes": {"conclusions": "a\x00b"}}).status_code == 422


def test_report_never_in_fox(admin, case):
    raw = admin.get(f"/api/cases/{case}/export").content
    doc = json.loads(raw)
    assert b"Ransomware intrusion" not in raw and b"EDR everywhere" not in raw and b"Initial access via VPN" not in raw
    assert not {"executive_summary", "conclusions", "lessons_learned"} & set(doc)
    assert doc["summary"] == "Client called on 2026-01-01. Scope: 3 sites."


def test_viewer_reads_but_cannot_edit_report(admin, case):
    uid = admin.post("/api/users", json={"username": "repviewer", "password": "a-long-password-123", "role": "user",
                                         "must_change_password": False}).json()["id"]
    assert admin.post(f"/api/cases/{case}/members", json={"user_id": uid, "permission": "viewer"}).status_code == 200
    v = login(TestClient(app), "repviewer", "a-long-password-123")
    assert v.get(f"/api/cases/{case}/report").status_code == 200
    assert v.patch(f"/api/cases/{case}/report", json={"changes": {"conclusions": "x"}}).status_code == 403


def test_excel_cover_and_executive_summary(admin, case):
    from openpyxl import load_workbook
    r = admin.get(f"/api/cases/{case}/export.xlsx")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames[:3] == ["Cover", "Report", "Timeline"]
    cover = {str(c.value) for row in wb["Cover"].iter_rows() for c in row if c.value is not None}
    assert "Aurora IR Web" in cover and "Incident report — REP-1" in cover and "Exported by" in cover
    assert "admin" in cover and any(re.match(r"\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2} UTC", v) for v in cover)
    assert wb["Cover"]._images, "logo on the cover"
    ex = {wb["Report"].cell(row=i, column=1).value: wb["Report"].cell(row=i, column=2).value
          for i in range(2, 6)}
    assert ex["Executive summary"].startswith("Ransomware intrusion.\n\nContained")
    assert ex["Lessons learned"] == "- MFA on VPN\n- EDR everywhere"
    assert "export_xlsx" == admin.get(f"/api/cases/{case}/audit?limit=1").json()[0]["action"]


def test_pdf_export(admin, case):
    r = admin.get(f"/api/cases/{case}/export.pdf")
    assert r.status_code == 200 and r.headers["content-type"] == "application/pdf"
    assert r.content.startswith(b"%PDF-") and r.headers["content-disposition"].endswith('.pdf"')
    from pypdf import PdfReader
    text = "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(r.content)).pages)
    for needle in ("Incident report", "REP-1", "Report", "Executive summary", "Lessons learned", "Charts", "Visual Timeline",
                   "Lateral Movement", "Activity Plot", "Timeline (", "Case Notes ("):
        assert needle in text, needle
    assert "<img src=/etc/passwd>" in text                  # hostile markup printed as text, not interpreted
    assert "root:x:0:0" not in text
    assert admin.get(f"/api/cases/{case}/audit?limit=1").json()[0]["action"] == "export_pdf"


def test_latex_export_escapes_everything(admin, case):
    r = admin.get(f"/api/cases/{case}/export.tex")
    assert r.status_code == 200 and r.headers["content-disposition"].endswith('.tex"')
    tex = r.content.decode()
    assert r"\begin{document}" in tex and r"\end{document}" in tex and "pgfplots" in tex
    assert r"\input{/etc/passwd}" not in tex and r"\write18{" not in tex.replace(r"\textbackslash{}write18", "")
    assert r"\textbackslash{}input\{/etc/passwd\}" in tex and r"50\%" in tex and r"\$x\^{}2\$" in tex
    body = tex[tex.index(r"\begin{document}"):]
    assert "Ransomware intrusion." in body and "EDR everywhere" in body
    # braces balanced (no group can be opened or closed by case data)
    stripped = re.sub(r"\\[{}]", "", tex)
    assert stripped.count("{") == stripped.count("}")


def test_api_key_can_export_all_formats(admin, case):
    secret = admin.post("/api/keys", json={"name": "exporter", "profile": "read"}).json()["secret"]
    c = TestClient(app)
    c.headers["Authorization"] = f"Bearer {secret}"
    for ext in ("", ".xlsx", ".pdf", ".tex"):
        assert c.get(f"/api/cases/{case}/export{ext}").status_code == 200, ext
    assert c.get(f"/api/cases/{case}/report").status_code == 200


def test_pdf_export_survives_cells_taller_than_a_page(admin):
    """A record field can hold 20000 characters, more than one table row can show on a page: ReportLab cannot
    split a row, so the PDF shortens such cells (full text stays in the Excel and .fox exports)."""
    fx = json.loads(json.dumps(FIXTURE))
    fx["case_id"] = "REP-LONG"
    words = ("lateral movement via psexec to the file server " * 450)[:20000]
    fx["timeline"][0]["event_data"] = words
    fx["timeline"][0]["notes"] = "C:\\" + "A" * 5000                # one unbroken token in a narrow column
    r = admin.post("/api/cases/import", content=json.dumps(fx))
    assert r.status_code == 200, r.text
    cid = r.json()["id"]
    r = admin.get(f"/api/cases/{cid}/export.pdf")
    assert r.status_code == 200, r.text[:300]
    from pypdf import PdfReader
    text = "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(r.content)).pages)
    assert "truncated" in text and "lateral movement via psexec" in text
    assert admin.get(f"/api/cases/{cid}/export.tex").status_code == 200
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(admin.get(f"/api/cases/{cid}/export.xlsx").content))
    assert any(c.value == words for row in wb["Timeline"].iter_rows() for c in row)


def test_engagement_end_platform_only_and_in_reports(admin, case):
    """Engagement end: set in Case Details, shown on the report covers, never written to or read from a .fox."""
    r = admin.patch(f"/api/cases/{case}", json={"changes": {"start_date": "2026-01-01", "end_date": "2026-02-14"}})
    assert r.status_code == 200, r.text
    assert admin.get(f"/api/cases/{case}").json()["meta"]["end_date"] == "2026-02-14"
    bad = admin.patch(f"/api/cases/{case}", json={"changes": {"end_date": "2025-12-31"}})
    assert bad.status_code == 422 and "before its start" in bad.text
    fox = admin.get(f"/api/cases/{case}/export").content
    assert b"end_date" not in fox and b"2026-02-14" not in fox
    from openpyxl import load_workbook
    cover = load_workbook(io.BytesIO(admin.get(f"/api/cases/{case}/export.xlsx").content))["Cover"]
    assert any(c.value == "2026-02-14" for row in cover.iter_rows() for c in row)
    from pypdf import PdfReader
    pdf = admin.get(f"/api/cases/{case}/export.pdf").content
    assert "Engagement end" in PdfReader(io.BytesIO(pdf)).pages[0].extract_text()
    assert "2026-02-14" in PdfReader(io.BytesIO(pdf)).pages[0].extract_text()
    assert "2026-02-14" in admin.get(f"/api/cases/{case}/export.tex").text
    # a replace from the (end-date-less) .fox keeps it; a file that carries one is ignored
    assert admin.post(f"/api/cases/{case}/replace", content=fox).status_code == 200
    assert admin.get(f"/api/cases/{case}").json()["meta"]["end_date"] == "2026-02-14"
    doc = json.loads(fox)
    doc.update(case_id="REP-END", end_date="2030-01-01")
    cid = admin.post("/api/cases/import", content=json.dumps(doc)).json()["id"]
    assert admin.get(f"/api/cases/{cid}").json()["meta"]["end_date"] == ""
