"""Default record order (UI/API/MCP/.fox/Excel), case and user lists, and single-line text."""
import io
import json

import pytest
from fastapi.testclient import TestClient

from tests.test_app import FIXTURE, login  # sets the test environment first
from tests.test_api import bearer, make_key
from app import cases as C
from app import schema as S
from app.db import cases_db
from app.main import app


def rec(recid, **data):
    return {"recid": recid, "w2ui": {}, **data}


@pytest.fixture(scope="module")
def admin():
    with TestClient(app) as c:
        yield login(c)


@pytest.fixture(scope="module")
def case(admin):
    """A case whose records are deliberately stored in a scrambled order."""
    fx = json.loads(json.dumps(FIXTURE))
    fx["case_id"] = "ORD-1"
    fx["systems"] = [rec(1, text="web10"), rec(2, text="Alpha"), rec(3, text="web2"), rec(4, text="beta")]
    fx["investigators"] = [rec(1, text="zz"), rec(2, text="AA"), rec(3, text="mm")]
    fx["timeline"] = [rec(1, date_time="2026-01-01 10:00:00", event_data="b"), rec(2, date_time="", event_data="none"),
                      rec(3, date_time="2026-03-01 00:00:00", event_data="c"),
                      rec(4, date_time="2025-12-31 23:59:59", event_data="a")]
    fx["investigated_systems"] = [rec(1, hostname="web2"), rec(2, hostname="Alpha"), rec(3, hostname="web10")]
    fx["malware"] = [rec(1, text="z.exe", hostname="web2"), rec(2, text="b.exe", hostname=""),
                     rec(3, text="a.exe", hostname="web2"), rec(4, text="m.exe", hostname="Alpha")]
    fx["compromised_accounts"] = [rec(1, account_name="bob", domain="CORP"), rec(2, account_name="adm", domain=""),
                                  rec(3, account_name="Alice", domain="corp"), rec(4, account_name="svc", domain="ACME")]
    fx["network_indicators"] = [rec(1, ip="10.0.0.10"), rec(2, ip="", domainname="evil.example"),
                                rec(3, ip="10.0.0.9"), rec(4, ip="2001:db8::1"), rec(5, ip="9.9.9.9")]
    fx["exfiltration"] = [rec(1, exfiltime="2026-01-01 00:00:00", filename="old"), rec(2, exfiltime="", filename="none"),
                          rec(3, exfiltime="2026-02-01 00:00:00", filename="new")]
    fx["osint"] = [rec(1, text="Shodan"), rec(2, text="abuse.ch"), rec(3, text="VT")]
    fx["evidence"] = [rec(1, date_acquired="2026-01-01 00:00:00", name="b"), rec(2, date_acquired="2026-02-01 00:00:00", name="z"),
                      rec(3, date_acquired="2026-01-01 00:00:00", name="a"), rec(4, date_acquired="", name="c")]
    fx["actions"] = [rec(1, date_added=300, task="third"), rec(2, date_added=100, task="first"),
                     rec(3, date_added=200, task="second")]
    fx["casenotes"] = [rec(1, date_added=100, note="old"), rec(2, date_added=300, note="new"),
                       rec(3, date_added=200, note="mid")]
    r = admin.post("/api/cases/import", content=json.dumps(fx))
    assert r.status_code == 200, r.text
    return r.json()["id"]


EXPECTED = {   # section -> recids in default order
    "timeline": [3, 1, 4, 2],                   # newest first, empty date last
    "investigated_systems": [2, 1, 3],          # Alpha, web2, web10 (natural order)
    "malware": [4, 3, 1, 2],                    # by host (Alpha, web2, empty last), then filename
    "compromised_accounts": [4, 3, 1, 2],       # ACME/svc, corp/Alice, CORP/bob, (no domain)/adm
    "network_indicators": [5, 3, 1, 4, 2],      # 9.9.9.9 < 10.0.0.9 < 10.0.0.10 < IPv6, no IP last
    "exfiltration": [3, 1, 2],                  # newest first, empty last
    "osint": [2, 1, 3],                         # abuse.ch, Shodan, VT
    "systems": [2, 4, 3, 1],                    # Alpha, beta, web2, web10 (natural, case-insensitive)
    "investigators": [2, 3, 1],                 # AA, mm, zz
    "evidence": [2, 3, 1, 4],                   # newest first, then name; empty date last
    "actions": [2, 3, 1],                       # oldest first
    "casenotes": [2, 3, 1],                     # newest first
}


def recids(admin, case, section, **params):
    r = admin.get(f"/api/cases/{case}/sections/{section}", params=params)
    assert r.status_code == 200, r.text
    return [x["recid"] for x in r.json()["records"]]


@pytest.mark.parametrize("section", list(EXPECTED))
def test_api_default_order(admin, case, section):
    assert recids(admin, case, section) == EXPECTED[section]


def test_schema_publishes_default_order(admin):
    secs = {s["key"]: s for s in admin.get("/api/schema").json()["sections"]}
    assert secs["timeline"]["default_order"] == [{"field": "date_time", "order": "desc"}]
    assert secs["malware"]["default_order"] == [{"field": "hostname", "order": "asc"}, {"field": "text", "order": "asc"}]
    assert all(secs[k]["default_order"] for k in EXPECTED)


def test_explicit_sort_keeps_empty_last(admin, case):
    assert recids(admin, case, "timeline", sort="date_time", order="asc") == [4, 1, 3, 2]
    assert recids(admin, case, "network_indicators", sort="ip", order="desc") == [4, 1, 3, 5, 2]
    assert recids(admin, case, "systems", sort="recid", order="desc") == [4, 3, 2, 1]
    assert admin.get(f"/api/cases/{case}/sections/systems", params={"sort": "nope"}).status_code == 422


def test_new_record_takes_its_place(admin, case):
    r = admin.post(f"/api/cases/{case}/sections/timeline", json={"data": {"date_time": "2026-02-01 00:00:00"}})
    assert r.status_code == 200, r.text
    new = r.json()["recid"]
    assert recids(admin, case, "timeline")[:2] == [3, new]
    admin.post(f"/api/cases/{case}/sections/timeline/{new}/delete", json={"version": r.json()["version"]})


def test_fox_and_excel_exports_use_default_order(admin, case):
    doc = json.loads(admin.get(f"/api/cases/{case}/export").content)
    for section, order in EXPECTED.items():
        assert [r["recid"] for r in doc[section]] == order, section
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(admin.get(f"/api/cases/{case}/export.xlsx").content))
    for section, order in EXPECTED.items():
        ws = wb[S.SECTIONS[section].label.replace("/", "-")]
        assert [ws.cell(row=i, column=1).value for i in range(2, ws.max_row + 1)] == order, section


def test_read_key_sees_default_order(admin, case):
    ro = bearer(make_key(admin, "order-reader", "read")[0])
    recs = ro.get(f"/api/cases/{case}/sections/timeline").json()["records"]
    assert [r["recid"] for r in recs] == EXPECTED["timeline"]


# ------------------------------------------------------------------ cases and users

def test_case_list_natural_desc_and_search(admin):
    ids = []
    for cid in ("IR-608", "IR-1000", "IR-609", "IR-99"):
        r = admin.post("/api/cases", json={"case_id": cid, "client": f"client {cid}"})
        assert r.status_code == 200, r.text
        ids.append(r.json()["id"])
    listed = [c["case_id"] for c in admin.get("/api/cases").json() if c["case_id"].startswith("IR-")]
    assert listed == ["IR-1000", "IR-609", "IR-608", "IR-99"]
    assert [c["case_id"] for c in admin.get("/api/cases", params={"q": "ir-60"}).json()] == ["IR-609", "IR-608"]
    ro = bearer(make_key(admin, "case-lister", "read")[0])
    assert [c["case_id"] for c in ro.get("/api/cases", params={"q": "IR-6"}).json()] == ["IR-609", "IR-608"]
    for pk, cid in zip(ids, ("IR-608", "IR-1000", "IR-609", "IR-99")):
        admin.post(f"/api/cases/{pk}/delete", json={"confirm": cid})


def test_user_list_ordered_and_searchable(admin):
    for name in ("zoe", "Bob", "alice"):
        assert admin.post("/api/users", json={"username": name, "password": "a-long-password-123",
                                               "role": "user"}).status_code == 200
    names = [u["username"] for u in admin.get("/api/users").json()]
    assert names == sorted(names, key=str.casefold)
    assert [u["username"] for u in admin.get("/api/users", params={"q": "BOB"}).json()] == ["Bob"]
    for u in admin.get("/api/users").json():
        if u["username"] in ("zoe", "Bob", "alice"):
            admin.delete(f"/api/users/{u['id']}")


# ------------------------------------------------------------------ single-line text

@pytest.mark.parametrize("text", ["two\nlines", "cr\rlf", "crlf\r\n", "sep x", "para x"])
def test_line_breaks_rejected_everywhere(admin, case, text):
    for section, field in (("timeline", "event_data"), ("casenotes", "note"), ("systems", "description"),
                           ("osint", "desc")):
        data = {field: text} | ({"note": text} if section == "casenotes" else {})
        r = admin.post(f"/api/cases/{case}/sections/{section}", json={"data": data})
        assert r.status_code == 422 and "line breaks" in r.text, (section, r.text)


def test_tab_and_unicode_still_allowed(admin, case):
    r = admin.post(f"/api/cases/{case}/sections/casenotes", json={"data": {"note": "tab\there · ünïcödé ✓"}})
    assert r.status_code == 200, r.text


def test_api_key_rejects_line_breaks(admin, case):
    api = bearer(make_key(admin, "nl-writer")[0])
    r = api.post(f"/api/cases/{case}/sections/casenotes", json={"data": {"note": "a\nb"}})
    assert r.status_code == 422 and "line breaks" in r.text


def test_import_joins_line_breaks(admin):
    fx = json.loads(json.dumps(FIXTURE))
    fx["case_id"] = "NL-1"
    fx["summary"] = "line one\r\nline two"
    fx["timeline"][0]["event_data"] = "first\n\n  second third"
    r = admin.post("/api/cases/import", content=json.dumps(fx))
    assert r.status_code == 200, r.text
    assert any("line breaks" in w for w in r.json()["warnings"])
    pk = r.json()["id"]
    assert admin.get(f"/api/cases/{pk}").json()["meta"]["summary"] == "line one\nline two"   # summary: multi-line
    ev = admin.get(f"/api/cases/{pk}/sections/timeline/1").json()["data"]["event_data"]
    assert ev == "first second third"


def test_summary_multiline_but_single_line_in_fox(admin, case):
    """Case summary (= executive-summary context) may span lines; the .fox gets it as one line."""
    r = admin.patch(f"/api/cases/{case}", json={"changes": {"summary": "para one\nline two\r\n\npara two"}})
    assert r.status_code == 200, r.text
    assert admin.get(f"/api/cases/{case}").json()["meta"]["summary"] == "para one\nline two\n\npara two"
    fox = json.loads(admin.get(f"/api/cases/{case}/export").content)
    assert fox["summary"] == "para one line two para two"
    assert admin.patch(f"/api/cases/{case}", json={"changes": {"summary": "bad\x00null"}}).status_code == 422


# ------------------------------------------------------------------ strict JSON (security review)

@pytest.mark.parametrize("bad", ['"x": NaN', '"x": Infinity', '"x": -Infinity', '"x": 1e999'])
def test_import_rejects_non_json_numbers(admin, bad):
    """Python's json accepts these; Aurora's JSON.parse does not, so such a file would export unreadable."""
    fx = json.loads(json.dumps(FIXTURE))
    fx["case_id"] = "NAN-1"
    raw = json.dumps(fx)[:-1] + ", " + bad + "}"
    r = admin.post("/api/cases/import", content=raw)
    assert r.status_code == 422 and "not valid JSON" in r.text, r.text


def test_admin_reset_and_deactivation_end_sessions(admin):
    """Containing a compromised account: the user's sessions end immediately."""
    assert admin.post("/api/users", json={"username": "tokuser", "password": "a-long-password-123",
                                           "role": "user", "must_change_password": False}).status_code == 200
    uid = next(u["id"] for u in admin.get("/api/users").json() if u["username"] == "tokuser")
    with TestClient(app) as c:
        user = login(c, "tokuser", "a-long-password-123")
        assert user.get("/api/me").status_code == 200
        assert admin.patch(f"/api/users/{uid}", json={"password": "another-long-pass-456"}).status_code == 200
        assert user.get("/api/me").status_code == 401
        user = login(c, "tokuser", "another-long-pass-456")
        admin.patch(f"/api/users/{uid}", json={"active": False})
        assert user.get("/api/me").status_code == 401
    admin.delete(f"/api/users/{uid}")
