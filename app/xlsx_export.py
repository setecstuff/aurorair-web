"""
Excel (.xlsx) export of a case: the "spreadsheet of doom".

A Cover sheet (application, logo, case presentation, exporter, UTC time, case metadata), a Report sheet (the
case's Report page), then one worksheet per section (same order as the sidebar, default record order). The data comes from
the same consistent snapshot that the .fox export uses (cases.build_fox), so the exports always agree.

Safety: case data contains attacker-controlled strings. Every text cell is written as an explicit string
(never a formula), so values such as `=cmd|' /C calc'!A0` or `=HYPERLINK(...)` cannot execute or phone
home when the workbook is opened (CSV/formula injection). Characters that are illegal in XML are dropped,
and cells longer than Excel's 32,767-character limit are truncated with a marker.
"""
from __future__ import annotations

import io
from datetime import date, datetime, timezone

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import schema as S

SHEET_ORDER = ["timeline", "investigated_systems", "malware", "compromised_accounts", "network_indicators",
               "exfiltration", "osint", "systems", "investigators", "evidence", "actions", "casenotes"]
EXCEL_MAX_CELL = 32767
DT_FORMAT = "yyyy-mm-dd hh:mm:ss"
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
TOP_WRAP = Alignment(vertical="top", wrap_text=True)
TOP = Alignment(vertical="top")


def _sheet_title(label: str, used: set) -> str:
    t = "".join("-" if c in '[]:*?/\\' else c for c in label)[:31].strip("'") or "Sheet"
    base, n = t, 2
    while t.lower() in used:
        t = f"{base[:28]} {n}"
        n += 1
    used.add(t.lower())
    return t


def _text(v) -> str:
    s = ILLEGAL_CHARACTERS_RE.sub("", str(v))
    if len(s) > EXCEL_MAX_CELL:
        marker = " …[truncated, see .fox]"
        s = s[:EXCEL_MAX_CELL - len(marker)] + marker
    return s


def _put(ws, row: int, col: int, value, number_format: str | None = None, wrap: bool = False):
    cell = ws.cell(row=row, column=col)
    if isinstance(value, str):
        cell.value = _text(value)
        cell.data_type = "s"            # explicit string: never interpreted as a formula
    else:
        cell.value = value
    if number_format:
        cell.number_format = number_format
    cell.alignment = TOP_WRAP if wrap else TOP
    return cell


def _typed(f: S.Field, v):
    """Convert a stored Aurora value to a native Excel value (UTC times stay naive = UTC)."""
    if v is None or v == "":
        return None, None
    try:
        if f.kind == "datetime":
            return datetime.strptime(v, "%Y-%m-%d %H:%M:%S"), DT_FORMAT
        if f.kind == "date":
            return datetime.strptime(v, "%Y-%m-%d").date(), "yyyy-mm-dd"
        if f.kind == "auto_ms" and isinstance(v, (int, float)) and not isinstance(v, bool):
            return datetime.fromtimestamp(v / 1000, tz=timezone.utc).replace(tzinfo=None), DT_FORMAT
    except (ValueError, OverflowError, OSError):
        return str(v), None
    if f.kind == "bool":
        return (v is True), None
    if f.kind == "int" and isinstance(v, int) and not isinstance(v, bool):
        return v, "0"
    return (v if isinstance(v, str) else str(v)), None


def _width(f: S.Field) -> float:
    w = f.width or ""
    if w.endswith("px"):
        return max(8, min(60, int(w[:-2]) / 7))
    if w.endswith("%"):
        return 60 if f.kind == "longtext" else 30
    return 18


def build_workbook(data: dict) -> bytes:
    """`data` comes from exports.collect(): the .fox document, the executive summary and export details."""
    from .exports import APP_NAME, metadata_rows, report_parts, section_counts
    from .security import LOGO_PNG
    doc = data["doc"]
    wb = Workbook()
    wb.properties.creator = APP_NAME
    wb.properties.title = f"{doc.get('case_id', '')} - case data"
    used: set = set()

    # Cover (secrets such as MISP/VirusTotal API keys are deliberately not exported)
    ws = wb.active
    ws.title = _sheet_title("Cover", used)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 90
    if LOGO_PNG.exists():
        from openpyxl.drawing.image import Image as XLImage
        img = XLImage(str(LOGO_PNG))
        img.width = img.height = 96
        ws.add_image(img, "B2")
    for r in range(2, 7):
        ws.row_dimensions[r].height = 15
    _put(ws, 8, 2, APP_NAME).font = Font(size=11, color="5B6675")
    _put(ws, 9, 2, f"Incident report — {doc.get('case_id', '')}").font = Font(size=20, bold=True, color="14304F")
    _put(ws, 10, 2, doc.get("client", "")).font = Font(size=14, bold=True)
    r = 12
    for k, v in metadata_rows(data):
        _put(ws, r, 2, k).font = Font(bold=True)
        _put(ws, r, 3, v, wrap=True)
        r += 1
    r += 1
    _put(ws, r, 2, "Records per section").font = Font(bold=True, color="14304F")
    r += 1
    for k, n in section_counts(data):
        _put(ws, r, 2, k)
        _put(ws, r, 3, n).alignment = Alignment(horizontal="left")
        r += 1
    r += 1
    _put(ws, r, 2, "Context").font = Font(bold=True, color="14304F")
    _put(ws, r + 1, 2, data["report"].get("context", ""), wrap=True)
    ws.merge_cells(start_row=r + 1, start_column=2, end_row=r + 1, end_column=3)
    ws.row_dimensions[r + 1].height = max(30, 15 * min(40, data["report"].get("context", "").count("\n") + 2))

    # Executive summary (never part of the .fox)
    ws = wb.create_sheet(_sheet_title("Report", used))
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 120
    for c, h in enumerate(("Part", "Text"), 1):
        hc = _put(ws, 1, c, h)
        hc.font, hc.fill = HEADER_FONT, HEADER_FILL
    for i, (label, text) in enumerate(report_parts(data), 2):
        _put(ws, i, 1, label).font = Font(bold=True)
        _put(ws, i, 2, text, wrap=True)
        ws.row_dimensions[i].height = max(30, min(409, 15 * (text.count("\n") + 1 + len(text) // 140)))
    ws.freeze_panes = "A2"

    for key in SHEET_ORDER:
        sec = S.SECTIONS[key]
        ws = wb.create_sheet(_sheet_title(sec.label, used))
        fields = [f for f in sec.fields]
        headers = ["ID"] + [f"{f.label} (UTC)" if f.kind in ("datetime", "auto_ms") else f.label for f in fields]
        for c, h in enumerate(headers, 1):
            hc = _put(ws, 1, c, h)
            hc.font, hc.fill = HEADER_FONT, HEADER_FILL
            hc.alignment = Alignment(vertical="center", wrap_text=True)
        for r, rec in enumerate(doc.get(key, []), 2):
            _put(ws, r, 1, rec.get("recid"))
            for c, f in enumerate(fields, 2):
                value, fmt = _typed(f, rec.get(f.name))
                if value is not None:
                    _put(ws, r, c, value, fmt, wrap=(f.kind == "longtext"))
        ws.column_dimensions["A"].width = 6
        for c, f in enumerate(fields, 2):
            ws.column_dimensions[get_column_letter(c)].width = _width(f)
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, len(doc.get(key, [])) + 1)}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
